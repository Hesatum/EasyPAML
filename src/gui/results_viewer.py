"""Results panel: summary, LRT tables, sites, branch analysis, export."""

import tkinter as tk
from tkinter import ttk

import customtkinter as ctk
from pathlib import Path
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('TkAgg')
from scipy import stats
import re
import threading
from html import escape as html_escape
from src.backend import lrt_stats
from src.backend.site_map import attach_original_positions
from src.backend.version import version_string
from . import charts
from .summary_tab import SummaryTab, ensure_tree_style, gene_verdict
from .branch_tab import BranchTab
from .gui_texts import TEXTS, tr
from .ui_helpers import (CURRENT_THEME, FONT_MONO, FONT_SIZE, FONT_UI, PALETTE, RADIUS, SPACE, fit_to_screen,
                         add_tooltip, LoadingOverlay, ask_save_file, open_folder, show_message)


class ResultsViewerWindow(SummaryTab, BranchTab, ctk.CTkToplevel):
    """Results panel window."""
    
    COLORS = {
        'bg_dark':        PALETTE['bg_window'],
        'bg_card':        PALETTE['bg_surface'],
        'bg_card_hover':  PALETTE['bg_elevated'],
        'bg_feed':        PALETTE['bg_panel'],
        'bg_sidebar':     PALETTE['bg_panel'],
        'bg_input':       PALETTE['bg_inset'],
        'text_primary':   PALETTE['text_primary'],
        'text_secondary': PALETTE['text_secondary'],
        'text_tertiary':  PALETTE['text_tertiary'],
        'text_muted':     PALETTE['text_muted'],
        'accent_blue':        PALETTE['accent_blue'],
        'accent_blue_hover':  PALETTE['accent_fill'],
        'accent_blue_light':  PALETTE['accent_text'],
        'accent_cyan':    PALETTE['accent_cyan'],
        'accent_cyan_hover': PALETTE['info_fill'],
        'accent_purple':  PALETTE['accent_purple'],
        'accent_purple_hover': '#7c3aed',
        'success':        PALETTE['success_text'],
        'success_hover':  PALETTE['success_fill'],
        'success_light':  PALETTE['success_light'],
        'warning':        PALETTE['warning_text'],
        'warning_hover':  PALETTE['warning_fill'],
        'danger':         PALETTE['danger_text'],
        'danger_hover':   PALETTE['danger_fill'],
        'info':           PALETTE['info_text'],
        'border':         PALETTE['divider'],
        'border_hover':   PALETTE['control_border_hover'],
    }

    # ── display helpers ─────────────────────────────

    @staticmethod
    def _font(size: str = 'sm', weight: str = 'normal'):
        return (FONT_UI, FONT_SIZE[size], weight)

    @staticmethod
    def _mono(size: str = 'sm', weight: str = 'normal'):
        return (FONT_MONO, FONT_SIZE[size], weight)

    @staticmethod
    def _style_combo(combo):
        combo.configure(fg_color=PALETTE['bg_inset'], border_color=PALETTE['control_border'],
                        button_color=PALETTE['control_border'],
                        button_hover_color=PALETTE['control_border_hover'],
                        dropdown_fg_color=PALETTE['bg_elevated'],
                        dropdown_hover_color=PALETTE['bg_elevated_hover'],
                        dropdown_text_color=PALETTE['text_primary'],
                        text_color=PALETTE['text_primary'], border_width=1,
                        corner_radius=RADIUS['field'],
                        font=(FONT_UI, FONT_SIZE['sm']), dropdown_font=(FONT_UI, FONT_SIZE['sm']))
        return combo

    def _style_tabs(self, tabs):
        """Mark the active tab."""
        seg = tabs._segmented_button
        seg.configure(font=(FONT_UI, FONT_SIZE['md'], 'bold'))

        def restyle(*_):
            current = tabs.get()
            self._ensure_tab(current)
            for name, btn in seg._buttons_dict.items():
                btn.configure(text_color='#ffffff' if name == current else PALETTE['text_secondary'])

        orig_set = tabs.set

        def set_and_restyle(name):
            orig_set(name)
            restyle()
        tabs.set = set_and_restyle
        tabs.configure(command=restyle)
        restyle()
    
    def __init__(self, parent, output_folder: Path):
        super().__init__(parent)
        self.title(TEXTS["viewer_window_title"])
        # never larger than the screen
        fit_to_screen(self, 1400, 900, min_w=1024, min_h=640)
        self.bind("<Escape>", lambda e: self.destroy())
        
        self.configure(fg_color=self.COLORS['bg_dark'])
        self.after(100, self.lift)
        self.after(150, self.focus_force)
        self.output_folder = output_folder
        self.df = None
        self.tag_columns = {}

        # the window shows at once; files are read in a thread behind a spinner and
        # the tabs are built when the data is ready
        self._overlay = LoadingOverlay(self, TEXTS["loading_results"])
        self._load_result = {}
        self._loader = threading.Thread(target=self._load_in_background, daemon=True)
        self._loader.start()
        self.after(50, self._wait_for_data)

    def _load_in_background(self):
        try:
            ok = self._load_data()
            if ok:
                self._extract_tag_columns()
                self._sort_by_significance()
                self._gene_verdicts()
            self._load_result['ok'] = ok
        except Exception as exc:
            self._load_result['error'] = exc

    def _wait_for_data(self):
        if self._loader.is_alive():
            self.after(50, self._wait_for_data)
            return
        if not self._load_result.get('ok'):
            self._overlay.close()
            self._show_error(TEXTS["viewer_error_no_tsv"] if 'error' not in self._load_result
                             else str(self._load_result['error']))
            return
        self.setup_ui()
        self.update_idletasks()
        self._overlay.frame.lift()
        self.after(10, self._overlay.close)
    
    def _load_data(self) -> bool:
        """Load analysis_summary.tsv, regenerating it when model folders are missing from it."""
        tsv_file = self.output_folder / "analysis_summary.tsv"
        if not tsv_file.exists():
            return False

        try:
            # ── 1. Detect model subfolders not yet in the TSV ──────────────────
            _legacy = {'BranchSite_A': 'Branch-site', 'BranchSite_A_null': 'Branch-site_null'}
            _exclude = {'reports'}
            found_models = set()
            for _item in self.output_folder.iterdir():
                if _item.is_dir() and _item.name not in _exclude:
                    found_models.add(_legacy.get(_item.name, _item.name))

            # Models already described in the TSV (read header only)
            import csv
            with open(tsv_file, newline='', encoding='utf-8', errors='ignore') as _f:
                _header = next(csv.reader(_f, delimiter='\t'), [])
            known_models = {col.rsplit('_', 1)[0] for col in _header
                            if col.endswith('_lnL')}

            missing_models = found_models - known_models
            if missing_models:
                print(f"[INFO] Models in subfolders missing from TSV: {missing_models}")
                print("[INFO] Regenerating analysis_summary.tsv to include all models...")
                from src.backend.codeml_backend import CodemlBatchAnalysis
                CodemlBatchAnalysis._regenerate_analysis_summary(self.output_folder)
                print("[OK] analysis_summary.tsv regenerated successfully.")

            # ── 2. Load (possibly updated) TSV ────────────────────────────────
            self.df = pd.read_csv(tsv_file, sep='\t')
            self.df = self.df.replace(['NA', 'nan', '', 'None'], np.nan)

            numeric_cols = [col for col in self.df.columns if col not in ('Gene', 'status', 'reason')]
            for col in numeric_cols:
                self.df[col] = pd.to_numeric(self.df[col], errors='coerce')

            self._recover_missing_omegas()

            # ── 3. orphan runs (.ctl without results), shown as a warning
            try:
                from src.backend.codeml_backend import CodemlBatchAnalysis
                self._orphaned_analyses = CodemlBatchAnalysis._find_orphaned_analyses(
                    self.output_folder
                )
            except Exception:
                self._orphaned_analyses = {}

            print(f"[OK] Dados carregados: {len(self.df)} genes")
            print(f"[OK] Colunas: {list(self.df.columns)}")
            return True
        except Exception as e:
            print(f"[ERR] Could not load the data: {e}")
            return False
    
    def _recover_missing_omegas(self):
        """Fill missing ω values from the result files."""
        from src.backend.sites_parser import SitesParser
        omega_cols = [col for col in self.df.columns if '_omega' in col]
        
        for omega_col in omega_cols:
            model_name = omega_col.replace('_omega', '')
            
            missing_rows = self.df[self.df[omega_col].isna()].index
            
            if len(missing_rows) == 0:
                continue
            
            print(f"\n[INFO] Recuperando omegas faltantes para {model_name}...")
            
            for idx in missing_rows:
                gene_name = self.df.loc[idx, 'Gene']
                
                results_file = self._find_results_file(gene_name, model_name)
                
                if results_file:
                    try:
                        omega = SitesParser.extract_omega_robust(results_file)
                        if omega is not None:
                            self.df.loc[idx, omega_col] = omega
                            print(f"  [OK] {gene_name} ({model_name}): w = {omega:.4f}")
                        else:
                            print(f"  [SKIP] {gene_name} ({model_name}): no value found")
                    except Exception as e:
                        print(f"  [ERR] {gene_name} ({model_name}): {e}")
                else:
                    print(f"  [INFO] {gene_name} ({model_name}): file not found")
    
    def _find_results_file(self, gene_name: str, model_name: str):
        """Result file of a gene and model, also under the legacy BranchSite_A name."""
        results_file = self.output_folder / model_name / f"{gene_name}_{model_name}_results.txt"
        if results_file.exists():
            return results_file
        
        if model_name == 'Branch-site':
            results_file = self.output_folder / 'BranchSite_A' / f"{gene_name}_BranchSite_A_results.txt"
            if results_file.exists():
                return results_file
        
        return None
    
    def _extract_tag_columns(self):
        """Columns of per-label ω values."""
        tag_pattern = r'(.+?)_([^_]+)_(omega|lnL)$'
        
        self.tag_columns = {}
        
        for col in self.df.columns:
            match = re.match(tag_pattern, col)
            if match:
                model, tag, metric = match.groups()
                
                if model not in self.tag_columns:
                    self.tag_columns[model] = {'tags': set(), 'omega': {}, 'lnL': {}}
                
                self.tag_columns[model]['tags'].add(tag)
                if metric == 'omega':
                    self.tag_columns[model]['omega'][tag] = col
                elif metric == 'lnL':
                    self.tag_columns[model]['lnL'][tag] = col
        
        print(f"Tags detected: {self.tag_columns}")
    
    @staticmethod
    def _fmt_pval(p: float) -> str:
        """p in readable scientific notation, never "0.00000000"."""
        try:
            return lrt_stats.format_p(float(p))
        except (TypeError, ValueError):
            return "NA"

    def _show_error(self, message: str):
        """Show an error screen."""
        error_frame = ctk.CTkFrame(self, fg_color=self.COLORS['bg_dark'])
        error_frame.pack(fill='both', expand=True, padx=20, pady=20)
        
        ctk.CTkLabel(error_frame, text="[X]", font=(FONT_UI, 48)).pack(pady=20)
        ctk.CTkLabel(error_frame, text=message,
                    font=(FONT_UI, 14, "bold"),
                    text_color=self.COLORS['danger']).pack(pady=10)
        ctk.CTkLabel(error_frame, text=TEXTS["viewer_error_run_analysis"],
                    font=(FONT_UI, 11),
                    text_color=self.COLORS['text_tertiary']).pack()
    
    def setup_ui(self):
        """Build the window."""

        header = ctk.CTkFrame(self, fg_color=PALETTE['bg_panel'], corner_radius=0, height=52)
        header.pack(fill='x', padx=0, pady=0)
        header.pack_propagate(False)

        left = ctk.CTkFrame(header, fg_color='transparent')
        left.pack(side="left", padx=SPACE['xl'], pady=0, fill='y')
        ctk.CTkLabel(left, text=TEXTS["viewer_header_title"], font=self._font('lg', 'bold'),
                     text_color=PALETTE['text_primary']).pack(side='left')
        ctk.CTkLabel(left, text=TEXTS["viewer_header_subtitle"], font=self._font('xs'),
                     text_color=PALETTE['text_tertiary']).pack(side='left', padx=(SPACE['md'], 0))

        right = ctk.CTkFrame(header, fg_color='transparent')
        right.pack(side="right", padx=SPACE['xl'], pady=SPACE['sm'], fill='y')
        ctk.CTkButton(right, text=TEXTS["viewer_btn_open_output"], height=32,
                      fg_color='transparent', border_width=1,
                      border_color=PALETTE['control_border'], text_color=PALETTE['text_primary'],
                      hover_color=PALETTE['bg_elevated'], corner_radius=RADIUS['field'],
                      font=self._font('sm', 'bold'),
                      command=lambda: open_folder(self.output_folder, self)).pack(side='right', padx=(SPACE['md'], 0))
        ctk.CTkButton(right, text=TEXTS["viewer_btn_recompute"], height=32,
                      fg_color='transparent', border_width=1,
                      border_color=PALETTE['control_border'], text_color=PALETTE['text_primary'],
                      hover_color=PALETTE['bg_elevated'], corner_radius=RADIUS['field'],
                      font=self._font('sm', 'bold'),
                      command=self._recompute_summaries).pack(side='right', padx=(SPACE['md'], 0))
        ctk.CTkLabel(right, text=TEXTS["viewer_genes_loaded"].format(n=len(self.df)),
                     font=self._font('sm'), text_color=PALETTE['text_secondary']).pack(side='right')

        stats_frame = ctk.CTkFrame(self, fg_color='transparent')
        stats_frame.pack(fill='x', padx=SPACE['lg'], pady=(SPACE['md'], 0))
        self._create_stats_panel(stats_frame)

        # ── warning banner: orphan runs ──────────────
        orphaned = getattr(self, '_orphaned_analyses', {})
        if orphaned:
            warn_frame = ctk.CTkFrame(self, fg_color=PALETTE['warning_subtle'],
                                      corner_radius=RADIUS['card'])
            warn_frame.pack(fill='x', padx=SPACE['lg'], pady=(SPACE['sm'], 0))

            n_genes  = len(orphaned)
            n_models = sum(len(v) for v in orphaned.values())
            genes_list = ', '.join(sorted(orphaned.keys())[:5])
            if n_genes > 5:
                genes_list += f' … (+{n_genes - 5})'

            warn_text = tr(
                f"{n_genes} gene(s) começaram mas não terminaram ({n_models} execução(ões) "
                f"com .ctl e sem resultado).\nAfetados: {genes_list}\n"
                f"Rode esses genes de novo; o motivo está em genes_status.tsv / batch_analysis_log.txt.",
                f"{n_genes} gene(s) started but never finished ({n_models} run(s) with a .ctl "
                f"and no result).\nAffected: {genes_list}\n"
                f"Re-run those genes; the reason is in genes_status.tsv / batch_analysis_log.txt."
            )
            ctk.CTkLabel(
                warn_frame,
                text=warn_text,
                font=self._font('sm'),
                text_color=PALETTE['warning_fg'],
                justify='left',
                anchor='w',
            ).pack(padx=SPACE['md'], pady=SPACE['sm'], anchor='w')

        tabs = ctk.CTkTabview(self, fg_color=PALETTE['bg_surface'],
                              segmented_button_fg_color=PALETTE['bg_panel'],
                              segmented_button_selected_color=PALETTE['accent_fill'],
                              segmented_button_selected_hover_color=PALETTE['accent_fill'],
                              segmented_button_unselected_color=PALETTE['bg_panel'],
                              segmented_button_unselected_hover_color=PALETTE['bg_elevated_hover'],
                              text_color='#ffffff',
                              corner_radius=RADIUS['panel'],
                              border_width=0)
        tabs.pack(fill='both', expand=True, padx=SPACE['lg'], pady=(SPACE['xs'], SPACE['md']))
        self.tabs = tabs
        
        tabs.add(TEXTS["viewer_tab_summary"])
        tabs.add(TEXTS["viewer_tab_sites"])

        branchsite_cols = [col for col in self.df.columns if 'Branch-site_class' in col]
        if branchsite_cols:
            tabs.add(TEXTS["viewer_tab_branchsite_classes"])

        has_branch = bool(self._branch_genes())
        if has_branch:
            tabs.add(TEXTS["viewer_tab_branch"])

        self._create_summary_tab(tabs.tab(TEXTS["viewer_tab_summary"]))
        # the other tabs are built the first time they are opened
        self._tab_builders = {TEXTS["viewer_tab_sites"]: self._create_sites_tab}

        if branchsite_cols:
            self._tab_builders[TEXTS["viewer_tab_branchsite_classes"]] = self._create_branchsite_class_tab

        if has_branch:
            self._tab_builders[TEXTS["viewer_tab_branch"]] = self._create_tree_tab
        self._style_tabs(tabs)

    def _ensure_tab(self, name: str) -> None:
        """Build a tab the first time it is shown."""
        build = self._tab_builders.pop(name, None)
        if build is not None:
            build(self.tabs.tab(name))
    
    # ── p and q per gene and pair (computed for older folders) ──

    def _pair_values(self, null: str, alt: str) -> dict:
        """{gene: (lrt, p, q)} for a model pair; None if the pair is absent."""
        cache = getattr(self, '_pq_cache', None)
        if cache is None:
            cache = self._pq_cache = {}
        if (null, alt) in cache:
            return cache[(null, alt)]
        lcol = lrt_stats.lrt_column(null, alt)
        if lcol not in self.df.columns:
            cache[(null, alt)] = None
            return None
        info = lrt_stats.PAIRS.get((null, alt), {'df': 1, 'boundary': False})
        pcol, qcol = lrt_stats.p_column(null, alt), lrt_stats.q_column(null, alt)
        genes, lrts, ps = [], [], []
        skipped_failed = False
        for _, row in self.df.iterrows():
            lrt = row.get(lcol)
            if pd.isna(lrt):
                continue
            if row.get('status') == 'failed':
                # older folders computed this LRT from a stopped codeml; leave it out
                skipped_failed = True
                continue
            p = row.get(pcol) if pcol in self.df.columns else np.nan
            if pd.isna(p) or p <= 0:   # p = 0 only comes from rounding: recompute
                df_ = info['df'] or 1
                p = lrt_stats.p_value(max(0.0, float(lrt)), df_, boundary=info['boundary'])
            genes.append(row['Gene']); lrts.append(float(lrt)); ps.append(float(p))
        if qcol in self.df.columns and not skipped_failed:
            qmap = dict(zip(self.df['Gene'], self.df[qcol]))
            qs = [qmap.get(g, np.nan) for g in genes]
            if any(pd.isna(q) or q <= 0 for q in qs):
                qs = lrt_stats.bh_qvalues(ps)
        else:
            qs = lrt_stats.bh_qvalues(ps)
        out = {g: (l, p, q) for g, l, p, q in zip(genes, lrts, ps, qs)}
        cache[(null, alt)] = out
        return out

    def _recompute_summaries(self) -> None:
        """Rebuild the summary files of this folder from the codeml outputs and reopen
        the panel."""
        from src.backend.codeml_backend import CodemlBatchAnalysis
        try:
            CodemlBatchAnalysis.regenerate_summary_files(self.output_folder)
        except Exception as exc:
            show_message(self, TEXTS["msg_error"], str(exc), 'error')
            return
        parent, folder = self.master, self.output_folder
        self.destroy()
        ResultsViewerWindow(parent, folder)

    def _sites_chart(self, parent, df_sites, results_file, gene, model, method) -> None:
        """Positions of the sites along the CDS (charts.draw_sites), above the table."""
        import json
        from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg as _Canvas
        from matplotlib.figure import Figure as _Figure
        if df_sites is None or df_sites.empty:
            return
        sitemap = Path(results_file).with_name(f"{gene}_{model}_sitemap.json")
        removed: list = []
        length = None
        try:
            sm = json.loads(sitemap.read_text(encoding='utf-8'))
            length = int(sm['n_codons_in_alignment'])
            removed = sorted(set(range(1, length + 1)) - set(sm.get('kept_codons', [])))
        except Exception:
            pass
        pos_col = 'position_original' if 'position_original' in df_sites.columns and \
            df_sites['position_original'].notna().all() else 'position'
        positions = df_sites[pos_col].astype(float)
        length = length or int(positions.max())
        marks = df_sites['significance'].fillna('') if 'significance' in df_sites.columns else \
            np.where(df_sites['pr_w_gt_1'] >= 0.99, '**', np.where(df_sites['pr_w_gt_1'] >= 0.95, '*', ''))
        marks = list(marks)
        n1, n2 = marks.count('*'), marks.count('**')
        legend = TEXTS["sites_chart_title"].format(model=model, method=method, n=n1 + n2, n1=n1, n2=n2)
        self._last_sites_chart = dict(gene=gene, model=model, method=method, positions=list(positions),
                                      probs=list(df_sites['pr_w_gt_1'].astype(float)), marks=marks,
                                      length=length, removed=removed, legend=legend)
        c = charts.DARK if CURRENT_THEME['mode'] == 'dark' else charts.LIGHT
        c = dict(c, bg=PALETTE['bg_panel'])
        fig = _Figure(figsize=(7.5, 3.4), facecolor=c['bg'])
        charts.draw_sites(fig, positions, df_sites['pr_w_gt_1'].astype(float), marks, length, removed, c,
                          title=legend)
        if removed:
            note = ctk.CTkLabel(parent, text=TEXTS["sites_chart_removed"].format(n=len(removed)),
                                font=self._font('xs'), anchor='w', justify='left', wraplength=600,
                                text_color=PALETTE['text_tertiary'])
            note.pack(side='bottom', fill='x', padx=SPACE['sm'])
            parent.bind('<Configure>', lambda e: note.configure(wraplength=max(200, e.width - 24)), add='+')
        canvas = _Canvas(fig, master=parent)
        widget = canvas.get_tk_widget()
        widget.configure(height=240, highlightthickness=0, bg=c['bg'])
        widget.pack(fill='both', expand=True)
        canvas.draw()

    _SITE_TESTS = {'M8': (('M8a', 'M8'), ('M7', 'M8')), 'M2a': (('M1a', 'M2a'),),
                   'Branch-site': (('Branch-site_null', 'Branch-site'),)}

    def _significant_for(self, model: str) -> set:
        """Genes with q < 0.05 in the main test of this model."""
        genes = set()
        # the first test that ran: M8 vs M8a before M8 vs M7, as in the Summary conclusion
        pair = next((p for p in self._SITE_TESTS.get(model, ()) if self._pair_values(*p)), None)
        for gene, (_, _, q) in (self._pair_values(*pair) or {}).items() if pair else ():
            if pd.notna(q) and q < 0.05:
                genes.add(gene)
        return genes

    def _sort_by_significance(self) -> None:
        """Order genes by their smallest q (then p) over the positive-selection tests,
        or over every test when there is none; failed genes go last."""
        pairs = self._positive_tests() or [pair for pair in lrt_stats.PAIRS
                                            if self._pair_values(*pair)]
        best = {}
        for pair in pairs:
            for gene, (_, p, q) in (self._pair_values(*pair) or {}).items():
                key = (q if pd.notna(q) else 1.0, p if pd.notna(p) else 1.0)
                best[gene] = min(best.get(gene, key), key)
        failed = self.df.get('status', pd.Series(index=self.df.index, dtype=object)) == 'failed'
        order = sorted(range(len(self.df)), key=lambda i: (
            bool(failed.iloc[i]), best.get(self.df['Gene'].iloc[i], (2.0, 2.0)), str(self.df['Gene'].iloc[i])))
        self.df = self.df.iloc[order].reset_index(drop=True)
        self._gene_rank = {g: i for i, g in enumerate(self.df['Gene'])}

    def _positive_tests(self):
        return [pair for pair in lrt_stats.POSITIVE_SELECTION_PAIRS if self._pair_values(*pair)]

    def _create_stats_panel(self, parent):
        """Summary cards: genes, models, significant genes per test (q < 0.05) and
        failed genes."""
        tests = []
        for null, alt in self._positive_tests():
            vals = self._pair_values(null, alt)
            n_sig = sum(1 for _, _, q in vals.values() if pd.notna(q) and q < 0.05)
            tests.append((f"{alt} vs {null}", n_sig, len(vals)))
        n_failed = int((self.df['status'] == 'failed').sum()) if 'status' in self.df.columns else 0

        value_font = (FONT_UI, 20, 'bold')

        def card(col, weight=0):
            parent.grid_columnconfigure(col, weight=weight)
            c = ctk.CTkFrame(parent, fg_color=PALETTE['bg_surface'], corner_radius=RADIUS['card'])
            c.grid(row=0, column=col, sticky='nsew',
                   padx=(0 if col == 0 else SPACE['xs'], 0 if col == 3 else SPACE['xs']))
            inner = ctk.CTkFrame(c, fg_color='transparent')
            inner.pack(fill='both', expand=True, padx=SPACE['lg'], pady=SPACE['sm'])
            return inner

        def value(box, number, label, color):
            ctk.CTkLabel(box, text=number, font=value_font, text_color=color).pack(side='left')
            ctk.CTkLabel(box, text=label, font=self._font('sm'), text_color=PALETTE['text_secondary']
                         ).pack(side='left', padx=(SPACE['sm'], SPACE['lg']), pady=(4, 0))

        value(card(0), str(len(self.df)), TEXTS["stats_total_genes"], PALETTE['text_primary'])
        value(card(1), self._count_models(), TEXTS["stats_models_run"], PALETTE['text_primary'])
        box = card(2, weight=1)
        ctk.CTkLabel(box, text=TEXTS["stats_sig_short"], font=self._font('sm'),
                     text_color=PALETTE['text_secondary']).pack(side='left', padx=(0, SPACE['md']), pady=(4, 0))
        for name, n_sig, total in tests:
            value(box, f"{n_sig}/{total}", name,
                  PALETTE['success_fg'] if n_sig else PALETTE['text_secondary'])
        if not tests:
            value(box, "—", "", PALETTE['text_secondary'])
        value(card(3), str(n_failed), TEXTS["stats_failed"],
              PALETTE['danger_fg'] if n_failed else PALETTE['text_secondary'])

    def _beb_sites(self, gene: str, model: str, threshold: float = 0.95):
        rf = self._find_results_file(gene, model)
        if not rf:
            return None
        try:
            from src.backend.sites_parser import SitesParser
            df = SitesParser.parse_sites_from_file(rf, method='BEB')
            if df.empty:
                df = SitesParser.parse_sites_from_file(rf, method='NEB')
            if df.empty:
                return df
            df = attach_original_positions(df, rf)
            return df[df['pr_w_gt_1'] >= threshold]
        except Exception:
            return None

    @staticmethod
    def _conclusion(sig_by_pair: dict, test_rows: list):
        """One plain sentence for a gene, from the q-values of its tests, and its colour."""
        kind, text = gene_verdict(sig_by_pair, test_rows)
        color = {'positive': PALETTE['success_fg'], 'none': PALETTE['text_secondary']}.get(kind, PALETTE['warning_fg'])
        return (text if kind in ('positive', 'none') else "⚠ " + text), color

    def _status_columns(self) -> dict:
        cache = getattr(self, '_status_cache', None)
        if cache is None:
            cache = self._status_cache = {'reason': {}, 'notes': {}}
            f = self.output_folder / 'genes_status.tsv'
            if f.exists():
                try:
                    st = pd.read_csv(f, sep='\t', dtype=str).fillna('')
                    for col in ('reason', 'notes'):
                        if col in st.columns:
                            cache[col].update(zip(st['Gene'], st[col]))
                except Exception:
                    pass
        return cache

    def _failure_reason(self, gene: str) -> str:
        return self._status_columns()['reason'].get(gene, '?')

    def _gene_notes(self, gene: str) -> str:
        """Warnings of a gene that ran (masked stop codon, excluded sequence)."""
        return self._status_columns()['notes'].get(gene, '')

    @staticmethod
    def _compact_reason(reason: str) -> str:
        """'M8: X; M8a: X; M7: Y' -> 'M8, M8a: X' and 'M7: Y'."""
        import re as _re
        parts = _re.split(r';\s+(?=(?:M\d\w*|Branch[-\w]*):\s)', str(reason).strip())
        groups: dict = {}
        for part in [x.strip() for x in parts if x.strip()]:
            m = _re.match(r'(M\d\w*|Branch[-\w]*):\s+(.*)$', part, _re.S)
            model, why = (m.group(1), m.group(2)) if m else ('', part)
            why = _re.sub(r';\s*(última linha|last line):\s*$', '', why.strip())
            groups.setdefault(why, []).append(model)
        return "\n".join((", ".join(x for x in ms if x) + ": " if any(ms) else "") + why
                         for why, ms in groups.items())

    def _create_sites_tab(self, parent):
        """Positive sites tab: BEB/NEB sites with alignment and codeml numbering."""
        ctrl = ctk.CTkFrame(parent, fg_color='transparent')
        ctrl.pack(fill='x', padx=SPACE['md'], pady=(SPACE['sm'], SPACE['xs']))
        line1 = ctk.CTkFrame(ctrl, fg_color='transparent')
        line1.pack(fill='x')
        lab = dict(font=self._font('sm', 'bold'), text_color=PALETTE['text_secondary'])

        models = [m for m in ('M8', 'M2a', 'Branch-site')
                  if (self.output_folder / m).exists() or
                  (m == 'Branch-site' and (self.output_folder / 'BranchSite_A').exists())] or ['M8']
        ctk.CTkLabel(line1, text=TEXTS["sites_label_model"], **lab).pack(side='left', padx=(0, SPACE['sm']))
        model_combo = self._style_combo(ctk.CTkComboBox(line1, values=models, width=130))
        model_combo.pack(side='left', padx=(0, SPACE['lg']))
        model_combo.set(models[0])
        ctk.CTkLabel(line1, text=TEXTS["sites_label_gene"], **lab).pack(side='left', padx=(0, SPACE['sm']))
        gene_combo = self._style_combo(ctk.CTkComboBox(line1, values=[], width=340))
        gene_combo.pack(side='left', padx=(0, SPACE['lg']))
        ctk.CTkLabel(line1, text=TEXTS["sites_label_analysis"], **lab).pack(side='left', padx=(0, SPACE['sm']))
        method_combo = self._style_combo(ctk.CTkComboBox(line1, values=['BEB', 'NEB'], width=90))
        method_combo.pack(side='left', padx=(0, SPACE['lg']))
        method_combo.set('BEB')
        filter_lbl = ctk.CTkLabel(line1, text=TEXTS["sites_label_filter"], **lab)
        filter_lbl.pack(side='left', padx=(0, SPACE['sm']))
        add_tooltip(filter_lbl, TEXTS["sites_legend"])
        p_filter = ctk.CTkEntry(line1, width=70, fg_color=PALETTE['bg_inset'], border_width=1,
                                border_color=PALETTE['control_border'], corner_radius=RADIUS['field'],
                                font=self._mono('sm'), text_color=PALETTE['text_primary'])
        p_filter.pack(side='left')
        p_filter.insert(0, "0.95")

        line2 = ctk.CTkFrame(ctrl, fg_color='transparent')
        line2.pack(fill='x', pady=(SPACE['xs'], 0))
        show_all = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(line2, text=TEXTS["sites_show_all"], variable=show_all, font=self._font('sm'),
                        text_color=PALETTE['text_secondary'], checkbox_width=18, checkbox_height=18,
                        command=lambda: update_gene_list()).pack(side='left')
        n_label = ctk.CTkLabel(line2, text="", font=self._font('xs'), text_color=PALETTE['text_tertiary'])
        n_label.pack(side='left', padx=(SPACE['sm'], SPACE['lg']))
        state = {'df': None, 'gene': '', 'model': ''}

        def _tsv(df):
            cols = ['position_original', 'position', 'amino_acid', 'pr_w_gt_1', 'significance',
                    'post_mean', 'post_se']
            out = df[[c for c in cols if c in df.columns]].rename(columns={
                'position_original': 'position_alignment', 'position': 'position_codeml'})
            return out.to_csv(sep='\t', index=False)

        def copy_sites():
            df = state['df']
            if df is None or df.empty:
                return
            self.clipboard_clear()
            self.clipboard_append(_tsv(df))
            show_message(self, TEXTS["msg_success"], TEXTS["sites_copied"].format(n=len(df)))

        def export_sites():
            df = state['df']
            if df is None or df.empty:
                return
            path = ask_save_file(self, TEXTS["sites_btn_export"], self.output_folder,
                                 initialfile=f"{state['gene']}_{state['model']}_sites.tsv",
                                 defaultextension='.tsv', filetypes=[('TSV', '*.tsv')])
            if path:
                Path(path).write_text(_tsv(df), encoding='utf-8')
                show_message(self, TEXTS["msg_success"], TEXTS["msg_exported_to"].format(path=path))

        def export_figure():
            args = getattr(self, '_last_sites_chart', None)
            if not args:
                return
            path = ask_save_file(self, TEXTS["sites_btn_figure"], self.output_folder,
                                 initialfile=f"{args['gene']}_{args['model']}_sites.png",
                                 defaultextension='.png', filetypes=[('PNG / PDF / SVG', '*.png *.pdf *.svg')])
            if not path:
                return
            try:
                from matplotlib.figure import Figure as _Figure
                c = charts.LIGHT
                fig = _Figure(figsize=(11, 4), facecolor=c['bg'])
                charts.draw_sites(fig, args['positions'], args['probs'], args['marks'], args['length'],
                                  args['removed'], c, title=f"{args['gene']} · {args['model']} · {args['method']}   "
                                  + args['legend'])
                fig.savefig(path, dpi=300, facecolor=c['bg'], bbox_inches='tight')
                show_message(self, TEXTS["msg_success"], TEXTS["msg_exported_to"].format(path=path))
            except Exception as e:
                show_message(self, TEXTS["msg_error"], TEXTS["msg_export_err"].format(error=e), 'error')

        for text, cmd in ((TEXTS["sites_btn_figure"], export_figure),
                          (TEXTS["sites_btn_export"], export_sites),
                          (TEXTS["sites_btn_copy"], copy_sites)):
            ctk.CTkButton(line2, text=text, command=cmd, height=28, fg_color='transparent',
                          border_width=1, border_color=PALETTE['control_border'],
                          text_color=PALETTE['text_primary'], hover_color=PALETTE['bg_elevated'],
                          corner_radius=RADIUS['field'],
                          font=self._font('sm', 'bold')).pack(side='right', padx=(SPACE['sm'], 0))

        head_host = ctk.CTkFrame(parent, fg_color='transparent')
        head_host.pack(fill='x', padx=SPACE['md'], pady=(SPACE['sm'], 0))
        table_frame = ctk.CTkFrame(parent, fg_color=PALETTE['bg_panel'], corner_radius=RADIUS['card'])
        table_frame.pack(fill='both', expand=True, padx=SPACE['md'], pady=(0, SPACE['md']))
        self._sites_head_host = head_host

        def update_gene_list(*args):
            model = model_combo.get()
            folder = self.output_folder / model
            if model == 'Branch-site' and not folder.exists():
                folder = self.output_folder / 'BranchSite_A'
            genes = []
            if folder.exists():
                rank = getattr(self, '_gene_rank', {})
                genes = sorted({m.group(1) for f in folder.glob('*_results.txt')
                                for m in [re.match(r'(.+?)_[A-Za-z0-9\-]+_results\.txt', f.name)] if m},
                               key=lambda g: (rank.get(g, len(rank)), g))
                if not show_all.get():
                    sig = self._significant_for(model)
                    genes = [g for g in genes if g in sig]
            n_label.configure(text=TEXTS["sites_genes_shown"].format(n=len(genes)) if show_all.get()
                              else TEXTS["sites_genes_significant"].format(n=len(genes)))
            gene_combo.configure(values=genes)
            gene_combo.set(genes[0] if genes else '')
            update_sites_table()

        def update_sites_table(*args):
            for w in table_frame.winfo_children() + head_host.winfo_children():
                w.destroy()
            self._last_sites_chart = None
            try:
                thr = float(p_filter.get().replace(',', '.'))
            except ValueError:
                thr = 0.95
            state['gene'], state['model'] = gene_combo.get(), model_combo.get()
            if not gene_combo.get():
                ctk.CTkLabel(table_frame, text=TEXTS["sites_no_significant_genes"].format(
                                 model=model_combo.get()),
                             font=self._font('md'), text_color=PALETTE['text_secondary'],
                             wraplength=900, justify='left').pack(pady=30, padx=SPACE['lg'])
                state['df'] = None
                return
            state['df'] = self._render_sites_table(table_frame, gene_combo.get(), model_combo.get(),
                                                   method_combo.get(), thr)

        def goto(gene, model):
            """Show a gene of a model (from the Summary)."""
            if model not in model_combo.cget('values'):
                return
            model_combo.set(model)
            if gene not in self._significant_for(model):
                show_all.set(True)
            update_gene_list()
            gene_combo.set(gene)
            update_sites_table()
            self.tabs.set(TEXTS["viewer_tab_sites"])
        self._sites_goto = goto

        model_combo.configure(command=update_gene_list)
        gene_combo.configure(command=update_sites_table)
        method_combo.configure(command=update_sites_table)
        p_filter.bind('<Return>', update_sites_table)
        p_filter.bind('<FocusOut>', update_sites_table)
        update_gene_list()

    def _render_sites_table(self, parent, gene_name: str, model_name: str, method: str,
                            p_threshold: float = 0.95):
        """Site table; returns the DataFrame shown, for copying and export."""
        results_file = self._find_results_file(gene_name, model_name) if gene_name else None
        if not results_file:
            ctk.CTkLabel(parent, text=TEXTS["sites_file_not_found"].format(
                             filename=f"{gene_name}_{model_name}_results.txt"),
                         font=(FONT_UI, 12), text_color=self.COLORS['warning']).pack(pady=50)
            return None
        try:
            from src.backend.sites_parser import SitesParser
            df_sites = SitesParser.parse_sites_from_file(results_file, method=method)
            df_sites = attach_original_positions(df_sites, results_file)
            df_f = SitesParser.filter_sites_by_pvalue(df_sites, p_threshold)
            if not df_f.empty:
                df_f = df_f.sort_values('position')
        except Exception as e:
            ctk.CTkLabel(parent, text=TEXTS["sites_parse_error"].format(error=str(e)),
                         font=(FONT_UI, 12), text_color=self.COLORS['danger']).pack(pady=50)
            return None

        from src.backend.sites_parser import SitesParser
        pc = SitesParser.extract_positive_class(results_file) or {}
        sub = f"{model_name} · {method}"
        if pc:
            sub += "  ·  " + TEXTS["summary_posclass"].format(w=f"{pc['omega']:.3f}", p1=f"{pc['p']:.3f}")
        host = getattr(self, '_sites_head_host', None)
        top = host if host is not None and host.winfo_exists() else parent
        banner = ctk.CTkFrame(top, fg_color='transparent')
        banner.pack(fill='x', pady=(0, SPACE['xs']))
        line = ctk.CTkFrame(banner, fg_color='transparent')
        line.pack(fill='x')
        ctk.CTkLabel(line, text=gene_name, font=self._font('md', 'bold'),
                     text_color=PALETTE['text_primary']).pack(side='left')
        ctk.CTkLabel(line, text=TEXTS['viewer_sites_count'].format(n=len(df_f)), font=self._font('md', 'bold'),
                     text_color=PALETTE['text_primary']).pack(side='left', padx=(SPACE['md'], 0))
        ctk.CTkLabel(line, text=sub, font=self._font('sm'),
                     text_color=PALETTE['text_secondary']).pack(side='left', padx=(SPACE['md'], 0))
        mapped = bool(len(df_f)) and bool(df_f.get('position_mapped', pd.Series([False])).all())
        if mapped or df_f.empty:
            ctk.CTkLabel(banner, text=TEXTS["sites_position_note"], font=self._font('xs'), wraplength=1150,
                         justify='left', anchor='w', text_color=PALETTE['text_tertiary']).pack(anchor='w')
        else:
            ctk.CTkLabel(banner, text="⚠ " + TEXTS["sites_unmapped_note"], font=self._font('xs'),
                         wraplength=1150, justify='left', anchor='w', text_color=PALETTE['warning_fg'],
                         fg_color=PALETTE['warning_subtle'], corner_radius=RADIUS['field'],
                         padx=SPACE['sm']).pack(anchor='w', fill='x', pady=(SPACE['xs'], 0))

        table_host = None
        if not df_f.empty:
            table_host = tk.Frame(parent, bg=PALETTE['bg_panel'])
            table_host.pack(side='right', fill='y', padx=(SPACE['sm'], SPACE['sm']), pady=SPACE['sm'])
        chart_host = ctk.CTkFrame(parent, fg_color='transparent')
        chart_host.pack(side='left', fill='both', expand=True, padx=(SPACE['sm'], 0), pady=SPACE['sm'])
        try:
            self._sites_chart(chart_host, df_sites, results_file, gene_name, model_name, method)
        except Exception as exc:
            print(f"[WARN] sites chart: {exc}")
        if df_f.empty:
            ctk.CTkLabel(top, text=TEXTS["sites_no_sites"].format(threshold=p_threshold),
                         font=self._font('md', 'bold'), text_color=PALETTE['text_primary'],
                         anchor='w').pack(fill='x', pady=(SPACE['sm'], 0))
            return df_f

        ensure_tree_style(self)
        cols = [('aln', 82, 'e'), ('codeml', 100, 'e'), ('aa', 40, 'center'), ('pr', 70, 'e'),
                ('sig', 44, 'center'), ('w', 120, 'e')]
        tree = ttk.Treeview(table_host, style='EP.Treeview', show='headings', selectmode='browse',
                            columns=[c[0] for c in cols])
        sb = ctk.CTkScrollbar(table_host, command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        sb.pack(side='right', fill='y')
        tree.pack(side='left', fill='y', expand=True)
        for (cid, w, anchor), title in zip(cols, TEXTS["sites_table_headers"]):
            tree.heading(cid, text=title, anchor=anchor)
            tree.column(cid, width=w, minwidth=w, anchor=anchor, stretch=False)
        tree.tag_configure('odd', background=PALETTE['row_alt'])
        tree.tag_configure('strong', foreground=PALETTE['success_fg'])
        for k, (_, row) in enumerate(df_f.iterrows()):
            pr = row['pr_w_gt_1']
            # codeml's own mark: it uses the unrounded value (0.990* is below 0.99)
            sig = row.get('significance')
            if not isinstance(sig, str):
                sig = '**' if pr >= 0.99 else ('*' if pr >= 0.95 else '')
            mean = row.get('post_mean', np.nan)
            se = row.get('post_se', np.nan)
            omega_txt = f"{mean:.3f} ± {se:.3f}" if pd.notna(mean) and pd.notna(se) else "—"
            cells = [str(int(row['position_original'])) if pd.notna(row.get('position_original')) else "?",
                     str(int(row['position'])), row.get('amino_acid', '?'), f"{pr:.3f}", sig, omega_txt]
            tags = (('odd',) if k % 2 else ()) + (('strong',) if sig == '**' else ())
            tree.insert('', 'end', values=cells, tags=tags)
        return df_f

    def _create_branchsite_class_tab(self, parent):
        """Branch-site class tab."""
        ctrl_frame = ctk.CTkFrame(parent, fg_color=self.COLORS['bg_card'],
                                 corner_radius=8, height=60)
        ctrl_frame.pack(fill='x', padx=10, pady=10)
        ctrl_frame.pack_propagate(False)
        
        ctk.CTkLabel(ctrl_frame, text=TEXTS["branchsite_classes_gene_label"], font=(FONT_UI, 11, "bold")).pack(side='left', padx=15, pady=10)
        
        genes = self.df['Gene'].tolist()
        gene_combo = ctk.CTkComboBox(ctrl_frame, values=genes, width=300)
        gene_combo.pack(side='left', padx=(0, 20))
        gene_combo.set(genes[0] if genes else "")
        
        table_frame = ctk.CTkScrollableFrame(parent, fg_color=self.COLORS['bg_feed'],
                                            corner_radius=8)
        table_frame.pack(fill='both', expand=True, padx=10, pady=10)
        
        def update_branchsite_table(*args):
            for widget in table_frame.winfo_children():
                widget.destroy()
            
            selected_gene = gene_combo.get()
            gene_row = self.df[self.df['Gene'] == selected_gene]
            
            if gene_row.empty:
                ctk.CTkLabel(table_frame, text=TEXTS["branchsite_classes_not_found"],
                           font=(FONT_UI, 11),
                           text_color=self.COLORS['warning']).pack(pady=50)
                return
            
            idx = gene_row.index[0]
            self._render_branchsite_class_table(table_frame, idx)
        
        gene_combo.configure(command=update_branchsite_table)
        update_branchsite_table()
    
    def _render_branchsite_class_table(self, parent, gene_idx: int):
        """Branch-site class table."""
        row = self.df.iloc[gene_idx]
        gene = row['Gene']
        
        header_frame = ctk.CTkFrame(parent, fg_color=self.COLORS['bg_card_hover'],
                                   corner_radius=8)
        header_frame.pack(fill='x', padx=8, pady=(8, 12))
        
        ctk.CTkLabel(header_frame, text=TEXTS["branchsite_classes_header"].format(gene=gene),
                    font=(FONT_UI, 12, "bold"),
                    text_color=self.COLORS['accent_cyan']).pack(pady=8)
        
        table_header_frame = ctk.CTkFrame(parent, fg_color=self.COLORS['bg_card_hover'],
                                         corner_radius=6)
        table_header_frame.pack(fill='x', padx=8, pady=(0, 4))
        
        bs_widths = [120, 150, 150, 150]
        headers = list(zip(TEXTS["branchsite_classes_table_headers"], bs_widths))

        for h_text, width in headers:
            ctk.CTkLabel(table_header_frame, text=h_text,
                        font=(FONT_UI, 11, "bold"),
                        text_color=self.COLORS['accent_blue_light'],
                        width=width).pack(side='left', padx=8, pady=8)
        
        for cls in ['0', '1', '2a', '2b']:
            class_col = f'Branch-site_class{cls}_fg_w'
            prop_col = f'Branch-site_class{cls}_prop'
            
            if class_col not in self.df.columns or prop_col not in self.df.columns:
                continue
            
            fg_w = row[class_col]
            prop = row[prop_col]
            
            
            row_frame = ctk.CTkFrame(parent, fg_color=self.COLORS['bg_card'],
                                    corner_radius=4, border_width=1,
                                    border_color=self.COLORS['bg_card_hover'])
            row_frame.pack(fill='x', padx=8, pady=2)
            
            if isinstance(prop, float):
                prop_str = f"{prop:.5f}"
            else:
                prop_str = "N/A"
            
            if isinstance(fg_w, float):
                fg_w_str = f"{fg_w:.5f}"
            else:
                fg_w_str = "N/A"
            
            bg_w_str = TEXTS["branchsite_classes_bg_w_placeholder"]
            
            cells = [
                (f"Class {cls}", 120),
                (prop_str, 150),
                (bg_w_str, 150),
                (fg_w_str, 150)
            ]
            
            for cell_text, width in cells:
                ctk.CTkLabel(row_frame, text=cell_text, font=(FONT_UI, 11),
                           text_color=self.COLORS['text_secondary'], width=width).pack(side='left', padx=8, pady=8)
        
        footer_frame = ctk.CTkFrame(parent, fg_color=self.COLORS['bg_card_hover'],
                                   corner_radius=6)
        footer_frame.pack(fill='x', padx=8, pady=(12, 8))
        
        ctk.CTkLabel(footer_frame, text=TEXTS["branchsite_classes_footer"],
                    font=(FONT_UI, 11),
                    text_color=self.COLORS['text_tertiary'],
                    wraplength=400).pack(pady=8, padx=8)
    
    def _detect_positive_selection(self) -> dict:
        """{gene: {test: {omega, p_value, q_value, lrt}}} for genes with q < 0.05 in at
        least one positive-selection test; omega is the positive class ω."""
        out = {}
        for null, alt in self._positive_tests():
            for gene, (lrt, p, q) in (self._pair_values(null, alt) or {}).items():
                if pd.notna(q) and q < 0.05:
                    row = self.df[self.df['Gene'] == gene].iloc[0]
                    out.setdefault(gene, {})[f"{alt} vs {null}"] = {
                        'omega': row.get(f'{alt}_w_pos', np.nan), 'p_value': p, 'q_value': q, 'lrt': lrt}
        return out

    def _n_failed(self) -> int:
        return int((self.df['status'] == 'failed').sum()) if 'status' in self.df.columns else 0

    def _count_models(self) -> str:
        """Number of models in the summary."""
        model_cols = [col for col in self.df.columns if '_lnL' in col or '_omega' in col]
        unique_models = set()
        
        for col in model_cols:
            model_name = col
            for suffix in ['_lnL', '_omega', '_np', '_time', '_stops']:
                if suffix in model_name:
                    model_name = model_name.replace(suffix, '')
                    break
            
            if model_name:
                unique_models.add(model_name)
        
        return str(len(unique_models))
    
    @staticmethod
    def _significant_rows(df_out: pd.DataFrame) -> pd.Series:
        """Rows marked significant in an export table (q < 0.05, or p < 0.05 when
        the test has no q)."""
        col = next((c for c in df_out.columns if c.startswith('significant (')), None)
        if col is None:
            return pd.Series(False, index=df_out.index)
        return df_out[col] == 'yes'

    @staticmethod
    def _significance_label(df_out: pd.DataFrame) -> str:
        return 'q < 0.05' if any(c == 'significant (q < 0.05)' for c in df_out.columns) else 'p < 0.05'

    def _build_export_df(self, lrt_col: str) -> pd.DataFrame:
        """One LRT comparison for export: genes where both models ran, with lnL and
        np of each model, ω values, 2Δℓ, df, p, q and BEB sites, unrounded."""
        lrt_series = pd.to_numeric(self.df[lrt_col], errors='coerce')
        mask = lrt_series.notna()
        if not mask.any():
            return pd.DataFrame()
        df_f = self.df[mask].copy()
        lrt_vals = lrt_series[mask]

        _info = lrt_stats.PAIRS.get(tuple(lrt_col.replace('lrt_', '').split('_vs_')), {})
        df_chi2 = 1 if _info.get('boundary') else (_info.get('df') or 1)

        # df = (np_Branch − np_M0) − (ntime_Branch − ntime_M0)
        branch_df_series = None
        if 'vs_Branch' in lrt_col and 'site' not in lrt_col.lower():
            if ('Branch_np' in df_f.columns and 'M0_np' in df_f.columns):
                raw_dfs = (
                    pd.to_numeric(df_f['Branch_np'], errors='coerce') -
                    pd.to_numeric(df_f['M0_np'],    errors='coerce')
                ).abs()
                if 'Branch_ntime' in df_f.columns and 'M0_ntime' in df_f.columns:
                    ntime_diff = (
                        pd.to_numeric(df_f['Branch_ntime'], errors='coerce') -
                        pd.to_numeric(df_f['M0_ntime'],    errors='coerce')
                    )
                    branch_df_series = (raw_dfs - ntime_diff).clip(lower=1)
                else:
                    branch_df_series = raw_dfs.clip(lower=1)

        if branch_df_series is not None:
            p_vals = pd.Series([
                float(stats.chi2.sf(x, df=int(d))) if pd.notna(x) and x > 0 else 1.0
                for x, d in zip(lrt_vals, branch_df_series)
            ], index=lrt_vals.index)
        else:
            p_vals = lrt_vals.apply(
                lambda x: float(stats.chi2.sf(x, df=df_chi2)) if pd.notna(x) and x > 0 else 1.0
            )

        out = pd.DataFrame({'Gene': df_f['Gene'].values})

        parts = lrt_col.replace('lrt_', '').split('_vs_')
        model_names = parts if len(parts) == 2 else []

        # ── ω columns ────────────────────────────────────────────────────
        # For M0 vs Branch: skip generic Branch_omega (replaced by per-tag below).
        # For all other models: add the generic omega column normally.
        _is_branch_lrt = (lrt_col == 'lrt_M0_vs_Branch')

        for mn in model_names:
            for col, label in ((f'{mn}_lnL', f'lnL ({mn})'), (f'{mn}_np', f'np ({mn})')):
                if col in df_f.columns:
                    out[label] = pd.to_numeric(df_f[col].values, errors='coerce')
        for mn in model_names:
            if _is_branch_lrt and mn == 'Branch':
                continue  # per-tag columns added below
            omega_col = f'{mn}_omega'
            if omega_col in df_f.columns:
                out[tr(f'ω médio ({mn})', f'mean ω ({mn})')] = pd.to_numeric(
                    df_f[omega_col].values, errors='coerce')
            for col, label in ((f'{mn}_w_pos', tr(f'ω classe positiva ({mn})', f'positive-class ω ({mn})')),
                               (f'{mn}_p_pos', tr(f'p₁ classe positiva ({mn})', f'positive-class p₁ ({mn})'))):
                if col in df_f.columns:
                    out[label] = pd.to_numeric(df_f[col].values, errors='coerce')

        # ── Per-tag ω columns for Branch model ───────────────────────────
        # PAML's "w (dN/dS) for branches:" lists groups as [bg, #1, #2, ...]
        # matching the user's branch labels in order of first appearance in tree.
        if _is_branch_lrt:
            from src.backend.sites_parser import SitesParser as _SP

            # 1) Try to get per-tag omegas from TSV columns (post-regeneration)
            _tag_cols_in_tsv = sorted(
                [c for c in df_f.columns
                 if re.match(r'^Branch_(background|#\d+)_omega$', c)],
                key=lambda c: (c != 'Branch_background_omega',
                               int(re.search(r'#(\d+)', c).group(1))
                               if re.search(r'#(\d+)', c) else 0)
            )
            if _tag_cols_in_tsv:
                for col in _tag_cols_in_tsv:
                    tag = col.replace('Branch_', '').replace('_omega', '')
                    label = 'ω (bg)' if tag == 'background' else f'ω ({tag})'
                    out[label] = pd.to_numeric(df_f[col].values, errors='coerce')
            else:
                # 2) Fallback: parse on-the-fly from result files
                _tag_data: dict = {}   # tag -> {gene -> omega}
                for gene in df_f['Gene'].values:
                    rf = self.output_folder / 'Branch' / f"{gene}_Branch_results.txt"
                    if not rf.exists():
                        continue
                    try:
                        for tag, omega in _SP.extract_omega_by_tags(rf).items():
                            _tag_data.setdefault(tag, {})[gene] = omega
                    except Exception:
                        pass

                # Sort: background first, then #1, #2, ... by number
                sorted_tags = sorted(
                    _tag_data.keys(),
                    key=lambda t: (t != 'background',
                                   int(re.search(r'(\d+)', t).group(1))
                                   if re.search(r'(\d+)', t) else 0)
                )
                for tag in sorted_tags:
                    label = 'ω (bg)' if tag == 'background' else f'ω ({tag})'
                    out[label] = [
                        float(_tag_data[tag][g]) if g in _tag_data[tag] else float('nan')
                        for g in df_f['Gene'].values
                    ]

        # ── Other per-tag ω columns (non-Branch models, e.g. Branch-site) ─
        tag_re = re.compile(r'^([A-Za-z0-9\-]+)_(.+)_omega$')
        for col in df_f.columns:
            m_col = tag_re.match(col)
            if m_col and col not in [f'{mn}_omega' for mn in model_names]:
                model_part, tag = m_col.group(1), m_col.group(2)
                # Skip Branch per-tag columns — already handled above
                if model_part == 'Branch' and re.match(r'^(background|#\d+)$', tag):
                    continue
                if any(mn.lower() == model_part.lower() for mn in model_names):
                    out[f'ω ({model_part}/{tag})'] = pd.to_numeric(df_f[col].values, errors='coerce')

        # same p and q as LRT_results.txt
        _pair = tuple(lrt_col.replace('lrt_', '').split('_vs_'))
        _pv = self._pair_values(*_pair) if len(_pair) == 2 and _pair in lrt_stats.PAIRS \
            and _pair[1] not in ('Branch',) else None
        q_vals = None
        if _pv:
            p_vals = pd.Series([_pv.get(g, (None, 1.0, None))[1] for g in df_f['Gene'].values],
                               index=lrt_vals.index)
            q_vals = pd.Series([_pv.get(g, (None, None, np.nan))[2] for g in df_f['Gene'].values],
                               index=lrt_vals.index)

        out['2Δℓ'] = lrt_vals.values
        # Branch: df = (np_Branch − np_M0) − (ntime_Branch − ntime_M0), per gene
        if branch_df_series is not None:
            out['df'] = branch_df_series.astype(int).values
        elif len(_pair) == 2 and _pair in lrt_stats.PAIRS and lrt_stats.PAIRS[_pair]['df']:
            out['df'] = lrt_stats.PAIRS[_pair]['df']
        out['p-value'] = p_vals.values
        if q_vals is not None:
            out['q-value (BH)'] = q_vals.values
            out['significant (q < 0.05)'] = ['yes' if pd.notna(q) and q < 0.05 else 'no'
                                             for q in q_vals.values]
        else:
            out['significant (p < 0.05)'] = ['yes' if p < 0.05 else 'no' for p in p_vals.values]

        # ── BEB positive sites (M2a and M8 only) ─────────────────────────
        # Format: "32 R* (8.200 ± 2.238); 91 G** (8.444 ± 1.804)"
        # Separator is ";" to avoid conflicts with CSV field delimiters.
        _SITES_MODELS = {'M1a_vs_M2a': 'M2a', 'M7_vs_M8': 'M8', 'M8a_vs_M8': 'M8'}
        _lrt_key = lrt_col.replace('lrt_', '')
        if _lrt_key in _SITES_MODELS:
            sites_model = _SITES_MODELS[_lrt_key]
            from src.backend.sites_parser import SitesParser as _SP2
            sites_col = []
            for gene in df_f['Gene'].values:
                rf = self.output_folder / sites_model / f"{gene}_{sites_model}_results.txt"
                if not rf.exists():
                    sites_col.append('')
                    continue
                try:
                    beb_df = _SP2.parse_sites_from_file(rf, method='BEB')
                    if beb_df.empty:
                        sites_col.append('')
                        continue
                    beb_df = attach_original_positions(beb_df, rf)
                    sig = beb_df[beb_df['pr_w_gt_1'] >= 0.95].sort_values('position')
                    if sig.empty:
                        sites_col.append('')
                        continue
                    parts_list = []
                    for _, sr in sig.iterrows():
                        star = sr['significance'] if sr['significance'] else (
                            '**' if sr['pr_w_gt_1'] >= 0.99 else '*')
                        # alignment numbering, with the codeml number in brackets when different
                        pos_o = int(sr['position_original']) if pd.notna(sr.get('position_original')) else int(sr['position'])
                        pos_c = int(sr['position'])
                        pos_txt = f"{pos_o}" if pos_o == pos_c else f"{pos_o} [codeml {pos_c}]"
                        parts_list.append(
                            f"{pos_txt} {sr['amino_acid']}{star} "
                            f"({sr['post_mean']:.3f} ± {sr['post_se']:.3f})"
                        )
                    sites_col.append('; '.join(parts_list))  # ";" avoids CSV conflicts
                except Exception:
                    sites_col.append('')
            out['Positive Sites (BEB, alignment numbering)'] = sites_col

        return out.reset_index(drop=True)

    def _export_excel(self):
        """Export to Excel, one sheet per test."""
        filepath = ask_save_file(self, TEXTS["dialog_save_as"], self.output_folder,
                                 initialfile="EasyPAML_results.xlsx", defaultextension=".xlsx",
                                 filetypes=[("Excel", "*.xlsx")])
        if not filepath:
            return
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            # fallback: plain pandas export (no formatting)
            lrt_cols = [c for c in self.df.columns if c.startswith('lrt_')]
            if not lrt_cols:
                show_message(self, TEXTS["msg_warning"], TEXTS["msg_no_lrt"], 'warning')
                return
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                for lrt_col in lrt_cols:
                    sheet_name = (lrt_col.replace('lrt_', '')
                                         .replace('_vs_', ' vs ')
                                         .replace('_', ' ')[:31])
                    df_out = self._build_export_df(lrt_col)
                    if not df_out.empty:
                        df_out.to_excel(writer, sheet_name=sheet_name, index=False)
            show_message(self, TEXTS["msg_success"], TEXTS["msg_exported_to"].format(path=filepath))
            return

        lrt_cols = [c for c in self.df.columns if c.startswith('lrt_')]
        if not lrt_cols:
            show_message(self, TEXTS["msg_warning"], TEXTS["msg_no_lrt"], 'warning')
            return

        # ── known sheet labels ──────────────────────────────────────────
        SHEET_LABELS = {
            'lrt_M0_vs_M1a':                       'M0 vs M1a',
            'lrt_M1a_vs_M2a':                      'M1a vs M2a',
            'lrt_M7_vs_M8':                        'M7 vs M8',
            'lrt_M0_vs_Branch':                    'M0 vs Branch',
            'lrt_Branch-site_null_vs_Branch-site': 'Branch-site',
            'lrt_M0_vs_Branch-site':               'M0 vs Branch-site',
        }

        # Style helpers — clean light-background professional theme
        HEADER_FILL   = PatternFill('solid', fgColor='1F3864')
        SIG05_FILL    = PatternFill('solid', fgColor='EBF5E0')
        ALT_FILL      = PatternFill('solid', fgColor='F5F7FB')
        HEADER_FONT   = Font(bold=True, color='FFFFFF', size=11)
        SIG05_FONT    = Font(color='2D6A1F', size=10)
        NORMAL_FONT   = Font(color='1A1A2E', size=10)
        CENTER        = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin          = Side(style='thin', color='CCCCCC')
        border        = Border(bottom=thin, left=thin, right=thin)

        try:
            with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
                sheets_written = 0
                for lrt_col in lrt_cols:
                    df_out = self._build_export_df(lrt_col)
                    if df_out.empty:
                        continue
                    sheet_name = SHEET_LABELS.get(lrt_col,
                                    lrt_col.replace('lrt_', '').replace('_vs_', ' vs ')
                                           .replace('_', ' ')[:31])
                    df_out.to_excel(writer, sheet_name=sheet_name, index=False)
                    ws = writer.sheets[sheet_name]

                    # ── style header row ──────────────────────────────
                    for cell in ws[1]:
                        cell.fill      = HEADER_FILL
                        cell.font      = HEADER_FONT
                        cell.alignment = CENTER

                    # ── style data rows ───────────────────────────────
                    sites_col_idx = None
                    for i, col in enumerate(df_out.columns, 1):
                        if col.startswith('Positive Sites (BEB'):
                            sites_col_idx = i
                    sig_rows = list(self._significant_rows(df_out))

                    for row_idx, row in enumerate(ws.iter_rows(min_row=2), 1):
                        if sig_rows[row_idx - 1]:
                            fill = SIG05_FILL
                            font = SIG05_FONT
                        else:
                            fill = ALT_FILL if row_idx % 2 == 0 else None
                            font = NORMAL_FONT
                        for col_i, cell in enumerate(row, 1):
                            cell.font      = font
                            cell.border    = border
                            # sites column: left-align, wrap
                            if col_i == sites_col_idx:
                                cell.alignment = Alignment(
                                    horizontal='left', vertical='top', wrap_text=True)
                            else:
                                cell.alignment = CENTER
                            if fill:
                                cell.fill = fill

                    # ── auto-fit column widths ────────────────────────
                    for col_cells in ws.columns:
                        header_val = col_cells[0].value or ''
                        # Sites column: fixed wide + row height
                        if str(header_val).startswith('Positive Sites (BEB'):
                            ws.column_dimensions[
                                get_column_letter(col_cells[0].column)
                            ].width = 60
                        else:
                            max_len = max(
                                len(str(c.value)) if c.value is not None else 0
                                for c in col_cells
                            )
                            ws.column_dimensions[
                                get_column_letter(col_cells[0].column)
                            ].width = min(max_len + 4, 35)

                    # Set row heights: header taller, data rows auto
                    ws.row_dimensions[1].height = 22
                    for r in range(2, ws.max_row + 1):
                        ws.row_dimensions[r].height = 18

                    sheets_written += 1

                # ── Summary sheet ─────────────────────────────────────
                summary_rows = []
                for lrt_col in lrt_cols:
                    df_out = self._build_export_df(lrt_col)
                    if df_out.empty:
                        continue
                    n_sig = int(self._significant_rows(df_out).sum())
                    sheet_name = SHEET_LABELS.get(lrt_col,
                                    lrt_col.replace('lrt_', '').replace('_vs_', ' vs ')
                                           .replace('_', ' '))
                    summary_rows.append({
                        'Test':           sheet_name,
                        'Genes':          len(df_out),
                        'Significant':    n_sig,
                        'Criterion':      self._significance_label(df_out),
                    })
                if summary_rows:
                    pd.DataFrame(summary_rows).to_excel(
                        writer, sheet_name='Summary', index=False)
                    ws_r = writer.sheets['Summary']
                    for cell in ws_r[1]:
                        cell.fill = HEADER_FILL
                        cell.font = HEADER_FONT
                        cell.alignment = CENTER
                    for col_cells in ws_r.columns:
                        max_len = max(
                            len(str(c.value)) if c.value is not None else 0
                            for c in col_cells
                        )
                        ws_r.column_dimensions[
                            get_column_letter(col_cells[0].column)
                        ].width = min(max_len + 4, 30)
                sheets_written = len(writer.book.sheetnames)

            show_message(self, TEXTS["msg_success"],
                         TEXTS["msg_excel_exported"].format(n=sheets_written, path=filepath))
        except Exception as e:
            show_message(self, TEXTS["msg_error"], TEXTS["msg_excel_err"].format(error=e), 'error')

    def _export_html(self):
        """Export an HTML report, one section per test."""
        filepath = ask_save_file(self, TEXTS["dialog_save_as"], self.output_folder,
                                 initialfile="EasyPAML_report.html", defaultextension=".html",
                                 filetypes=[("HTML", "*.html")])
        if not filepath:
            return

        SHEET_LABELS = {
            'lrt_M0_vs_M1a':                       ('M0 → M1a',  'Nearly neutral pre-test (M1a vs M0)'),
            'lrt_M1a_vs_M2a':                      ('M1a → M2a', 'Positive sites (M2a vs M1a)'),
            'lrt_M7_vs_M8':                        ('M7 → M8',   'Beta + ω > 1 (M8 vs M7)'),
            'lrt_M8a_vs_M8':                       ('M8a → M8',  'ω > 1 beyond neutral sites (M8 vs M8a)'),
            'lrt_M0_vs_Branch':                    ('M0 → Branch','Free branches (Branch vs M0)'),
            'lrt_Branch-site_null_vs_Branch-site': ('Branch-site','Episodic selection on the foreground'),
            'lrt_M0_vs_Branch-site':               ('M0 → Branch-site','Episodic selection (alt.)'),
        }

        try:
            positive_genes = self._detect_positive_selection()
            lrt_cols = [c for c in self.df.columns if c.startswith('lrt_')]
            now_str  = pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')

            # ── Build per-model section HTML ──────────────────────────
            sections_html = ''
            for lrt_col in lrt_cols:
                df_out = self._build_export_df(lrt_col)
                if df_out.empty:
                    continue
                df_out = df_out.drop(columns=[c for c in df_out.columns if c.startswith(('Positive Sites', 'np ('))])
                label, subtitle = SHEET_LABELS.get(lrt_col, (lrt_col, ''))

                th_cells = ''.join(f'<th>{c}</th>' for c in df_out.columns)
                tr_rows = ''
                sig_rows = list(self._significant_rows(df_out))
                for i_row, (_, row) in enumerate(df_out.iterrows()):
                    sig = sig_rows[i_row]
                    row_cls = ' class="sig"' if sig else ''
                    cells = ''
                    for col_name, val in row.items():
                        cell_cls = ''
                        if 'ω' in col_name and isinstance(val, float) and val > 1.0:
                            cell_cls = ' class="pos"'
                        if pd.isna(val):
                            display = '—'
                        elif isinstance(val, float) and col_name.startswith(('p-value', 'q-value')):
                            display = lrt_stats.format_p(val)
                        elif isinstance(val, float) and col_name.startswith(('lnL', '2Δℓ')):
                            display = f'{val:.3f}'
                        elif isinstance(val, float):
                            display = f'{val:.5g}'
                        else:
                            display = html_escape(str(val))
                        cells += f'<td{cell_cls}>{display}</td>'
                    tr_rows += f'<tr{row_cls}>{cells}</tr>\n'

                n_sig = int(self._significant_rows(df_out).sum())

                sections_html += f"""
        <section>
          <h2>{label}</h2>
          <p class="subtitle">{subtitle}</p>
          <p class="meta">{len(df_out)} gene(s) tested &nbsp;·&nbsp; {n_sig} significant ({self._significance_label(df_out).replace('<', '&lt;')})</p>
          <div style="overflow-x:auto">
          <table>
            <thead><tr>{th_cells}</tr></thead>
            <tbody>{tr_rows}</tbody>
          </table>
          </div>
        </section>
"""

            # ── Positive-selection cards ──────────────────────────────
            pos_html = ''
            verdicts = self._gene_verdicts()
            if positive_genes:
                for gene, signals in positive_genes.items():
                    kind = verdicts.get(gene, ('',))[0]
                    evidence = TEXTS["conclusion_short"].get(kind, '')
                    notes = self._gene_notes(gene)
                    sites_html = ''
                    model = 'M8' if any(t.startswith('M8 ') for t in signals) else 'M2a'
                    sites = self._beb_sites(gene, model)
                    if sites is not None and not sites.empty:
                        items = []
                        for _, sr in sites.sort_values('position').iterrows():
                            pos = sr.get('position_original')
                            pos = int(pos) if pd.notna(pos) else int(sr['position'])
                            star = sr['significance'] if isinstance(sr['significance'], str) and sr['significance'] \
                                else ('**' if sr['pr_w_gt_1'] >= 0.99 else '*')
                            items.append(f'<span class="site">{pos}&nbsp;{html_escape(str(sr["amino_acid"]))}{star}</span>')
                        sites_html = (f'<div class="sites"><b>{len(items)} {model} BEB site(s), Pr(ω&gt;1) ≥ 0.95'
                                      f'</b> (alignment numbering; ** ≥ 0.99): ' + ' '.join(items) + '</div>')
                    signals_inner = ''.join(
                        f'<div class="signal">{st}: p = {self._fmt_pval(sd["p_value"])}, '
                        f'q = {self._fmt_pval(sd["q_value"])}'
                        + (f', ω (positive class) = {sd["omega"]:.3f}' if pd.notna(sd["omega"]) else '')
                        + '</div>'
                        for st, sd in signals.items()
                    )
                    pos_html += (f'<div class="gene-card{"" if kind == "positive" else " weak"}">'
                                 f'<div class="gene-name">{html_escape(str(gene))}'
                                 + (f' <span class="evidence">{html_escape(evidence)}</span>' if evidence else '')
                                 + '</div>'
                                 f'{signals_inner}{sites_html}'
                                 + (f'<div class="note">⚠ {html_escape(notes.replace(" | ", "; "))}</div>' if notes else '')
                                 + '</div>\n')
            else:
                pos_html = '<p class="meta">No gene with a significant LRT (q &lt; 0.05)</p>' 

            html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>EasyPAML report</title>
<style>
*{{margin:0;padding:0;box-sizing:border-box}}
:root{{--bg:#f4f5f7;--card:#ffffff;--text:#1f2430;--muted:#555d6b;--line:#dde1e7;--head:#e8ebf0;
       --accent:#4338ca;--sig-bg:#e6f4ea;--sig:#14532d;--pos:#0e7490;--warn:#92400e}}
@media (prefers-color-scheme: dark){{:root{{--bg:#0d0d11;--card:#16161c;--text:#eeeef2;--muted:#a2a2b6;
       --line:#2a2a33;--head:#20202a;--accent:#818cf8;--sig-bg:#0b2016;--sig:#6ee7b7;--pos:#22d3ee;--warn:#fbbf24}}}}
body{{font-family:'Segoe UI',Roboto,'DejaVu Sans',sans-serif;background:var(--bg);color:var(--text);padding:32px 16px}}
.container{{max-width:1500px;margin:0 auto;background:var(--card);border-radius:12px;padding:36px;
            border:1px solid var(--line)}}
h1{{color:var(--accent);font-size:26px;margin-bottom:4px}}
h2{{font-size:19px;margin:36px 0 8px;padding-bottom:6px;border-bottom:1px solid var(--line)}}
.subtitle{{color:var(--muted);font-size:13px;margin-bottom:4px}}
.meta{{color:var(--muted);font-size:13px;margin-bottom:12px}}
.stats-grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin:16px 0}}
.stat-card{{padding:14px 16px;border-radius:8px;border:1px solid var(--line)}}
.stat-label{{font-size:12px;color:var(--muted);margin-bottom:4px}}
.stat-value{{font-size:26px;font-weight:700}}
table{{width:100%;border-collapse:collapse;margin:8px 0;font-size:13px}}
th{{background:var(--head);padding:8px 10px;text-align:left;font-weight:600;white-space:nowrap}}
td{{padding:7px 10px;border-bottom:1px solid var(--line);vertical-align:top;white-space:nowrap;
    font-variant-numeric:tabular-nums}}
tr.sig td{{background:var(--sig-bg);color:var(--sig)}}
td.pos{{color:var(--pos);font-weight:700}}
.gene-card{{border:1px solid var(--line);border-left:4px solid var(--sig);border-radius:8px;padding:12px 16px;margin:8px 0}}
.gene-name{{font-size:15px;font-weight:700;margin-bottom:6px}}
.signal{{font-size:13px;color:var(--muted);margin:2px 0}}
.evidence{{font-size:12px;font-weight:600;color:var(--sig);margin-left:8px}}
.gene-card.weak{{border-left-color:var(--warn)}}
.gene-card.weak .evidence{{color:var(--warn)}}
.sites{{font-size:13px;margin-top:8px;line-height:1.9}}
.site{{display:inline-block;padding:0 6px;margin-right:4px;border:1px solid var(--line);border-radius:4px;
       font-family:'DejaVu Sans Mono',Consolas,monospace;font-size:12px}}
.note{{font-size:13px;color:var(--warn);margin-top:6px}}
section{{margin-bottom:40px}}
.footer{{margin-top:40px;padding-top:16px;border-top:1px solid var(--line);color:var(--muted);font-size:12px}}
</style>
</head>
<body>
<div class="container">
  <h1>EasyPAML report</h1>
  <p class="meta">Generated {now_str} · {html_escape(version_string())}</p>

  <h2 style="margin-top:28px">Summary</h2>
  <div class="stats-grid">
    <div class="stat-card"><div class="stat-label">Genes</div>
      <div class="stat-value">{len(self.df)}</div></div>
    <div class="stat-card"><div class="stat-label">Models</div>
      <div class="stat-value">{self._count_models()}</div></div>
    <div class="stat-card"><div class="stat-label">Significant LRT (q &lt; 0.05)</div>
      <div class="stat-value">{len(positive_genes)}</div></div>
    <div class="stat-card"><div class="stat-label">Failed genes</div>
      <div class="stat-value">{self._n_failed()}</div></div>
  </div>

  <h2>Genes with a significant positive-selection LRT (q &lt; 0.05)</h2>
  {pos_html}

  <h2>Results by test</h2>
  {sections_html}

  <div class="footer">
    <p>Generated by EasyPAML with PAML/codeml. Methods and parameters: methods_text.txt and run_config.json in the results folder.</p>
  </div>
</div>
</body>
</html>
"""
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(html_content)
            show_message(self, TEXTS["msg_success"], TEXTS["msg_html_exported"].format(path=filepath))
        except Exception as e:
            show_message(self, TEXTS["msg_error"], TEXTS["msg_html_err"].format(error=e), 'error')